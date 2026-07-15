import sys, os, json
from datetime import datetime, timezone
sys.path.insert(0, os.path.dirname(__file__))

import streamlit as st
from dual_model_reviewer import _dispatch, FIRST_MAX_TOKENS, REVIEW_MAX_TOKENS, REVIEW_SYSTEM

ARCHIVE_FILE = os.path.join(os.path.dirname(__file__), "archives.json")

CLAUDE_MODELS = {
    "Opus 4.8  (most capable)": "claude-opus-4-8",
    "Sonnet 4.6  (balanced)":   "claude-sonnet-4-6",
    "Haiku 4.5  (fastest)":     "claude-haiku-4-5",
}
GEMINI_MODELS = {
    "Gemini 2.5 Flash  (default)": "gemini-2.5-flash",
    "Gemini 2.5 Pro  (capable)":   "gemini-2.5-pro",
    "Gemini 3.5 Flash  (latest)":  "gemini-3.5-flash",
}

def load_archives():
    if not os.path.exists(ARCHIVE_FILE):
        return []
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            with open(ARCHIVE_FILE, encoding=enc) as f:
                return json.load(f)
        except (UnicodeDecodeError, UnicodeError):
            continue
        except Exception:
            return []
    return []

def save_archive(entry):
    archives = load_archives()
    archives.insert(0, entry)
    with open(ARCHIVE_FILE, "w", encoding="utf-8") as f:
        json.dump(archives, f, indent=2, ensure_ascii=False)

def generate_pdf(turns: list, total_cost: float = 0.0) -> bytes:
    import re
    from fpdf import FPDF

    # Strip characters Segoe UI TTF doesn't cover.
    # Raw strings ensure \u escapes reach the regex engine, not Python's string parser.
    _UNSUPPORTED = re.compile(
        r'[\u2600-\u27FF'
        r'\uFE00-\uFE0F'
        r'\U0001F000-\U0001FFFF]',
        flags=re.UNICODE
    )

    def strip_md(text: str) -> str:
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'\*(.*?)\*', r'\1', text)
        text = re.sub(r'`{1,3}(.*?)`{1,3}', r'\1', text, flags=re.DOTALL)
        text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
        text = re.sub(r'^\s*[-*+]\s+', '• ', text, flags=re.MULTILINE)
        text = _UNSUPPORTED.sub('', text)
        return text.strip()

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    fonts = "C:\\Windows\\Fonts\\"
    pdf.add_font("Segoe",  fname=fonts + "segoeui.ttf")
    pdf.add_font("Segoe",  style="B", fname=fonts + "segoeuib.ttf")
    pdf.add_font("Segoe",  style="I", fname=fonts + "segoeuii.ttf")

    # Header
    pdf.set_font("Segoe", "B", 15)
    pdf.cell(0, 10, "Dual Model Reviewer", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Segoe", "I", 9)
    pdf.set_text_color(130, 130, 130)
    pdf.cell(0, 5, f"Exported {datetime.now().strftime('%Y-%m-%d %H:%M')}",
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(6)

    for msg in turns:
        if msg["role"] == "user":
            pdf.set_fill_color(232, 242, 255)
            pdf.set_font("Segoe", "B", 10)
            pdf.cell(0, 7, "You", new_x="LMARGIN", new_y="NEXT", fill=True)
            pdf.set_font("Segoe", "", 10)
            pdf.multi_cell(0, 5.5, strip_md(msg["content"]))
        else:
            pdf.set_fill_color(232, 255, 242)
            pdf.set_font("Segoe", "B", 10)
            pdf.cell(0, 7, "Assistant", new_x="LMARGIN", new_y="NEXT", fill=True)
            pdf.set_font("Segoe", "", 10)
            pdf.multi_cell(0, 5.5, strip_md(msg["content"]))
            if msg.get("initial") and msg["initial"] != msg["content"]:
                pdf.ln(1)
                pdf.set_font("Segoe", "I", 9)
                pdf.set_text_color(110, 110, 110)
                pdf.cell(0, 5, f"Initial response ({msg.get('first_label', '')})",
                         new_x="LMARGIN", new_y="NEXT")
                pdf.multi_cell(0, 5, strip_md(msg["initial"]))
                pdf.set_text_color(0, 0, 0)
        pdf.ln(4)

    if total_cost > 0:
        pdf.set_font("Segoe", "I", 9)
        pdf.set_text_color(130, 130, 130)
        pdf.cell(0, 5, f"Estimated session cost: ${total_cost:.4f}",
                 new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())

def render_cost_metrics(cost_info: dict):
    mc1, mc2, mc3 = st.columns(3)
    mc1.metric(f"Call 1 · {cost_info['first_label']}", f"${cost_info['c1_cost']:.4f}",
               f"{cost_info['c1_in']:,} in / {cost_info['c1_out']:,} out", delta_color="off")
    mc2.metric(f"Call 2 · {cost_info['second_label']}", f"${cost_info['c2_cost']:.4f}",
               f"{cost_info['c2_in']:,} in / {cost_info['c2_out']:,} out", delta_color="off")
    mc3.metric("Turn cost", f"${cost_info['c1_cost'] + cost_info['c2_cost']:.4f}")

# ── Page config + session state ───────────────────────────────────────────────
st.set_page_config(page_title="Dual Model Reviewer", page_icon="🔍", layout="wide")

for key, default in [
    ("first_history", []),   # neutral [{role, content}] sent to first model each turn
    ("display", []),         # [{role, content, initial?, cost_info?}] for UI
    ("total_cost", 0.0),
    ("upload_key", 0),       # increment to reset file uploader after send
]:
    if key not in st.session_state:
        st.session_state[key] = default

# ── Sidebar ───────────────────────────────────────────────────────────────────
with st.sidebar:
    st.header("Settings")

    claude_label = st.selectbox("Claude model", list(CLAUDE_MODELS))
    gemini_label = st.selectbox("Gemini model", list(GEMINI_MODELS))
    claude_model_id = CLAUDE_MODELS[claude_label]
    gemini_model_id = GEMINI_MODELS[gemini_label]

    first_model = st.radio(
        "Which model answers first?",
        options=["claude", "gemini"],
        format_func=lambda x: "🟣 Claude" if x == "claude" else "🔵 Gemini",
        horizontal=True,
    )
    second_model = "gemini" if first_model == "claude" else "claude"
    model_ids = {"claude": claude_model_id, "gemini": gemini_model_id}
    label_map  = {
        "claude": claude_label.split("(")[0].strip(),
        "gemini": gemini_label.split("(")[0].strip(),
    }

    st.divider()

    if st.session_state.total_cost > 0:
        st.metric("Session cost", f"${st.session_state.total_cost:.4f}")

    btn_col1, btn_col2 = st.columns(2)
    if btn_col1.button("🗑 New chat", use_container_width=True):
        st.session_state.first_history = []
        st.session_state.display = []
        st.session_state.total_cost = 0.0
        st.session_state.upload_key += 1
        st.rerun()

    if btn_col2.button("📥 Archive", use_container_width=True,
                       disabled=not st.session_state.display):
        save_archive({
            "timestamp":   datetime.now(timezone.utc).isoformat(),
            "first_model": first_model,
            "model_ids":   model_ids,
            "turns":       st.session_state.display,
            "total_cost":  st.session_state.total_cost,
        })
        st.success("Conversation archived.")

    if st.session_state.display:
        try:
            pdf_data = generate_pdf(st.session_state.display, st.session_state.total_cost)
            ts_str = datetime.now().strftime("%Y%m%d_%H%M")
            st.download_button("📄 Export PDF", data=pdf_data,
                               file_name=f"conversation_{ts_str}.pdf",
                               mime="application/pdf", use_container_width=True)
        except Exception as e:
            st.caption(f"PDF error: {e}")

    st.divider()
    st.header("Archives")
    archives = load_archives()
    if not archives:
        st.caption("No archives yet.")
    for i, entry in enumerate(archives):
        ts         = entry.get("timestamp", "")[:16].replace("T", " ")
        turns      = entry.get("turns", [])
        first_user = next((t["content"] for t in turns if t["role"] == "user"), "—")
        label      = f"{ts} · {first_user[:32]}{'…' if len(first_user) > 32 else ''}"
        with st.expander(label):
            for turn in turns:
                if turn["role"] == "user":
                    st.markdown(f"**You:** {turn['content']}")
                else:
                    st.markdown(turn["content"])
                    if turn.get("initial") and turn["initial"] != turn["content"]:
                        with st.expander("Initial response"):
                            st.markdown(turn["initial"])
            if entry.get("total_cost"):
                st.caption(f"Session cost: ${entry['total_cost']:.4f}")
            arc_col1, arc_col2 = st.columns(2)
            try:
                arc_pdf = generate_pdf(turns, entry.get("total_cost", 0.0))
                ts_slug = entry.get("timestamp", "")[:10]
                arc_col2.download_button("📄 PDF", data=arc_pdf,
                                         file_name=f"conversation_{ts_slug}.pdf",
                                         mime="application/pdf",
                                         key=f"pdf_{i}", use_container_width=True)
            except Exception:
                pass
            if arc_col1.button("📂 Load", key=f"load_{i}", use_container_width=True):
                st.session_state.display = turns
                # Rebuild first model's history using its original responses
                st.session_state.first_history = [
                    {"role": t["role"] if t["role"] == "user" else "assistant",
                     "content": t["content"] if t["role"] == "user" else t.get("initial", t["content"])}
                    for t in turns
                ]
                st.session_state.total_cost = entry.get("total_cost", 0.0)
                st.session_state.upload_key += 1
                st.rerun()

# ── Main ──────────────────────────────────────────────────────────────────────
missing = [k for k in ("ANTHROPIC_API_KEY", "GEMINI_API_KEY") if not os.getenv(k)]
if missing:
    st.warning(f"Missing env vars: {', '.join(missing)}")

# File uploader — resets after each send via upload_key
uploaded = st.file_uploader(
    "Attach a file to your next message (optional)",
    type=["jpg", "jpeg", "png", "gif", "webp", "pdf", "txt", "docx"],
    key=f"uploader_{st.session_state.upload_key}",
)

image_bytes = None
image_mime  = "image/jpeg"
pdf_bytes   = None
doc_text    = None

if uploaded:
    file_bytes = uploaded.read()
    fname = uploaded.name.lower()

    if uploaded.type.startswith("image/"):
        image_bytes = file_bytes
        image_mime  = uploaded.type
        st.image(image_bytes, width=240, caption=uploaded.name)

    elif fname.endswith(".pdf") or uploaded.type == "application/pdf":
        pdf_bytes = file_bytes
        st.info(f"📄 {uploaded.name} attached")

    elif fname.endswith(".txt") or uploaded.type == "text/plain":
        doc_text = file_bytes.decode("utf-8", errors="replace")
        st.info(f"📝 {uploaded.name} attached ({len(doc_text):,} chars)")

    elif fname.endswith(".docx"):
        try:
            import io
            from docx import Document as DocxDocument
            doc_text = "\n".join(
                p.text for p in DocxDocument(io.BytesIO(file_bytes)).paragraphs if p.text.strip()
            )
            st.info(f"📝 {uploaded.name} attached ({len(doc_text):,} chars)")
        except ImportError:
            st.error("python-docx not installed. Run: pip install python-docx")

# ── Render conversation history ───────────────────────────────────────────────
for msg in st.session_state.display:
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])
        if msg.get("initial") and msg["initial"] != msg["content"]:
            with st.expander(f"Initial · {msg.get('first_label', '')}"):
                st.markdown(msg["initial"])
        if "cost_info" in msg:
            render_cost_metrics(msg["cost_info"])

# ── Chat input ────────────────────────────────────────────────────────────────
if prompt := st.chat_input("Ask anything…", disabled=bool(missing)):

    # Prepend extracted text from TXT/DOCX into the prompt
    effective_prompt = f"[Attached document]\n{doc_text}\n\n{prompt}" if doc_text else prompt

    # Show user message immediately
    with st.chat_message("user"):
        st.markdown(prompt)
        if image_bytes:
            st.image(image_bytes, width=240)

    st.session_state.display.append({"role": "user", "content": prompt})

    # Call 1: first model with full history
    with st.chat_message("assistant"):
        with st.spinner(f"[1/2] {label_map[first_model]} generating…"):
            try:
                r1 = _dispatch(
                    first_model, effective_prompt,
                    max_tokens=FIRST_MAX_TOKENS,
                    model_id=model_ids[first_model],
                    image_bytes=image_bytes,
                    image_mime=image_mime,
                    pdf_bytes=pdf_bytes,
                    history=st.session_state.first_history,
                )
            except Exception as e:
                st.error(f"Error on call 1: {e}")
                st.stop()

        # Update first model's history with its own response
        st.session_state.first_history.append({"role": "user",      "content": prompt})
        st.session_state.first_history.append({"role": "assistant",  "content": r1["text"]})

        # Call 2: reviewer is always stateless (no history), PDF not re-sent
        review_prompt = f"Original question: {prompt[:200]}\n\nResponse to review:\n{r1['text']}"
        with st.spinner(f"[2/2] {label_map[second_model]} reviewing…"):
            try:
                r2 = _dispatch(
                    second_model, review_prompt,
                    system=REVIEW_SYSTEM,
                    max_tokens=REVIEW_MAX_TOKENS,
                    model_id=model_ids[second_model],
                    image_bytes=image_bytes,
                    image_mime=image_mime,
                    thinking=True,
                )
            except Exception as e:
                st.error(f"Error on call 2: {e}")
                st.stop()

        st.markdown(r2["text"])

        if r1["text"] != r2["text"]:
            with st.expander(f"Initial · {label_map[first_model]}"):
                st.markdown(r1["text"])

        turn_cost = r1["cost_usd"] + r2["cost_usd"]
        st.session_state.total_cost += turn_cost

        cost_info = {
            "first_label":  label_map[first_model],
            "second_label": label_map[second_model],
            "c1_cost": r1["cost_usd"], "c1_in": r1["input_tokens"], "c1_out": r1["output_tokens"],
            "c2_cost": r2["cost_usd"], "c2_in": r2["input_tokens"], "c2_out": r2["output_tokens"],
        }
        render_cost_metrics(cost_info)
        st.caption(f"Session total: ${st.session_state.total_cost:.4f}")

    st.session_state.display.append({
        "role":        "assistant",
        "content":     r2["text"],
        "initial":     r1["text"],
        "first_label": label_map[first_model],
        "cost_info":   cost_info,
    })

    # Reset file uploader so image doesn't re-attach to next message
    st.session_state.upload_key += 1
    st.rerun()
